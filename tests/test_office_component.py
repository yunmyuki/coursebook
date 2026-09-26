"""Actual LibreOffice PPT/PPTX conversion, without model calls or a system install."""
import tempfile,unittest
from pathlib import Path
from unittest.mock import patch
from pptx import Presentation
from pptx.util import Inches
from course_compiler.extract import extract_pptx,convert_office

class OfficeComponentTests(unittest.TestCase):
    def test_group_positions_table_and_speaker_notes_and_legacy_ppt(self):
        converter=Path('build/office-component/program/soffice.exe').resolve()
        if not converter.exists():self.skipTest('Bundled Office component not built')
        with tempfile.TemporaryDirectory(dir=Path('tmp').resolve()) as d:
            d=Path(d);prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6]);group=slide.shapes.add_group_shape();box=group.shapes.add_textbox(Inches(1),Inches(1),Inches(2),Inches(1));box.text='A grouped source paragraph';group.left=Inches(3);group.width=Inches(4)
            table=slide.shapes.add_table(2,2,Inches(1),Inches(4),Inches(4),Inches(1)).table;table.cell(0,0).text='Definition';table.cell(0,1).text='Value';table.cell(1,0).text='Gain';table.cell(1,1).text='+5.0%'
            slide.notes_slide.notes_text_frame.text='Speaker notes must remain complete.';source=d/'lecture.pptx';prs.save(source)
            with patch('course_compiler.extract.locate_converter',return_value=str(converter)):
                pages=extract_pptx(source,source,d/'site','lecture-office');self.assertEqual(len(pages),1)
                units=pages[0]['units'];self.assertTrue(any(u['type']=='speaker-note' and 'remain complete' in u['sourceText'] for u in units));self.assertEqual(sum(u['type']=='table-cell' for u in units),4)
                grouped=next(u for u in units if 'grouped' in u['sourceText']);self.assertAlmostEqual(grouped['position'][0],.3,places=2)
                legacy=convert_office(source,d/'legacy','ppt');self.assertTrue(legacy.exists());roundtrip=convert_office(legacy,d/'roundtrip','pptx');self.assertTrue(roundtrip.exists())

if __name__=='__main__':unittest.main()
