import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from pypdf import PdfWriter
from pypdf.generic import NameObject,DictionaryObject,DecodedStreamObject
from pptx import Presentation
from pptx.util import Inches
from course_compiler.extract import extract_pdf,extract_pptx,course_order_key


def make_pdf(path,count=1):
    writer=PdfWriter()
    for n in range(count):
        page=writer.add_blank_page(width=300,height=300)
        font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
        page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):writer._add_object(font)})})
        stream=DecodedStreamObject();stream.set_data(b'BT /F1 14 Tf 20 260 Td (Preserved definition) Tj 0 -35 Td (1. First bullet) Tj 0 -22 Td (2. Repeated definition) Tj ET')
        page[NameObject('/Contents')]=writer._add_object(stream)
    writer.write(path)


class ExtractionTests(unittest.TestCase):
    def test_course_order_prefers_filename_then_metadata(self):
        with tempfile.TemporaryDirectory(dir=Path('tmp').resolve()) as d:
            root=Path(d)
            files=[root/'Lecture10.pptx',root/'appendix.pptx',root/'02-introduction.pptx',root/'z-notes.pptx']
            for p in files:
                prs=Presentation();prs.core_properties.title='Week 3' if p.name=='z-notes.pptx' else 'Lecture 99'
                if p.name=='appendix.pptx':prs.core_properties.title='Appendix'
                prs.save(p)
            self.assertEqual([p.name for p in sorted(files,key=course_order_key)],['02-introduction.pptx','z-notes.pptx','Lecture10.pptx','appendix.pptx'])

    def test_pdf_text_and_independent_page_image(self):
        with tempfile.TemporaryDirectory(dir=Path('tmp').resolve()) as d:
            root=Path(d);source=root/'lecture.pdf';make_pdf(source)
            pages=extract_pdf(source,root/'site','lecture')
            self.assertEqual(len(pages),1);p=pages[0]
            for text in ('Preserved definition','First bullet','Repeated definition'):self.assertIn(text,p['rawText'])
            self.assertTrue((root/'site'/p['image']).exists())
            self.assertEqual(sum(u['type']=='bullet' for u in p['units']),2)
            self.assertTrue(all(len(u['position'])==4 for u in p['units']))

    def test_pptx_xml_tables_bullets_and_speaker_notes(self):
        with tempfile.TemporaryDirectory(dir=Path('tmp').resolve()) as d:
            root=Path(d);source=root/'lecture.pptx';rendered=root/'rendered.pdf';make_pdf(rendered,2)
            prs=Presentation()
            for number in (1,2):
                slide=prs.slides.add_slide(prs.slide_layouts[5]);slide.shapes.title.text='Slide '+str(number)
                tf=slide.shapes.add_textbox(Inches(1),Inches(1),Inches(5),Inches(1)).text_frame
                tf.text='Parent bullet';child=tf.add_paragraph();child.text='Child detail';child.level=1
                table=slide.shapes.add_table(2,2,Inches(1),Inches(3),Inches(5),Inches(1)).table
                table.cell(0,0).text='Definition';table.cell(0,1).text='Value';table.cell(1,0).text='Repeated';table.cell(1,1).text='0.01'
                slide.notes_slide.notes_text_frame.text='Speaker note with an example.'
            prs.save(source)
            # Converter integration is external; fixture lets XML preservation be checked independently.
            with patch('course_compiler.extract.convert_office',return_value=rendered):
                pages=extract_pptx(source,source,root/'site','lecture')
            self.assertEqual(len(pages),2)
            tids=[]
            for p in pages:
                self.assertTrue(any(u['type']=='speaker-note' and 'example' in u['sourceText'] for u in p['units']))
                self.assertTrue(any(u['type']=='bullet' and u['level']==1 for u in p['units']))
                cells=[u for u in p['units'] if u['type']=='table-cell'];self.assertEqual(len(cells),4)
                self.assertEqual(p['source']['file'],'lecture.pptx');tids.append(cells[0]['tableId'])
            self.assertNotEqual(*tids)

if __name__=='__main__':unittest.main()
