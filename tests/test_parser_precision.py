"""Regression cases for source crop fidelity, multi-column text and physical tables."""
import copy
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock,patch
from PIL import Image
from reportlab.pdfgen import canvas
from pypdf import PdfReader,PdfWriter
from pptx import Presentation
from pptx.util import Inches

from course_compiler.extract import extract_pdf,extract_pptx
from course_compiler.visual_input import render_region
from course_compiler.figures import add_figure
from course_compiler.hybrid_parser import parse,native_blocks
from course_compiler.parsers import commit_layout,table_cells
from course_compiler.pipeline_v2 import compile_course,course_id,fingerprint
from course_compiler.storage import atomic_json
from test_v2 import page,PROFILE


def text_line(text,box,font='Body',size=12):
    return {'text':text,'position':box,'typography':{'fontName':font,'fontSize':size,'dominance':1}}


class ParserPrecisionTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory(dir='tmp');self.root=Path(self.temp.name)
        (self.root/'sources').mkdir();self.source=self.root/'sources/lecture-test.pdf'
    def tearDown(self):self.temp.cleanup()

    def ruled_table(self):
        c=canvas.Canvas(str(self.source),pagesize=(600,400))
        c.rect(50,150,500,150);c.line(50,250,550,250);c.line(180,150,180,250)
        c.setFont('Helvetica',10);c.drawString(65,275,'Total compensation')
        c.drawString(65,200,'-6.25%');c.drawString(200,200,'0.01')
        c.save()
        return extract_pdf(self.source,self.root,'lecture-test')[0]

    def test_small_region_is_rendered_from_source_and_cache_is_reused(self):
        c=canvas.Canvas(str(self.source),pagesize=(2000,1000));c.setFont('Helvetica',5)
        c.drawString(100,800,'Small original text');c.save()
        p=extract_pdf(self.source,self.root,'lecture-test')[0]
        b=p['textLines'][0]['position']
        target,meta=render_region(p,self.root,b,self.root/'crops',padding=.002)
        self.assertEqual(meta['method'],'source-pdf-region');self.assertEqual(meta['contentBBox'],b)
        with Image.open(self.root/p['image']) as preview:
            self.assertGreater(meta['pixelSize'][1],(b[3]-b[1])*preview.height*3)
        self.assertLessEqual(max(meta['pixelSize']),4096);self.assertLessEqual(meta['pixelSize'][0]*meta['pixelSize'][1],6000000)
        stamp=target.stat().st_mtime_ns
        self.assertEqual(render_region(p,self.root,b,self.root/'crops',padding=.002)[0],target)
        self.assertEqual(target.stat().st_mtime_ns,stamp)

    def test_figure_uses_source_resolution_and_stable_anchor(self):
        c=canvas.Canvas(str(self.source),pagesize=(1200,600));c.rect(100,100,300,200);c.save()
        p=extract_pdf(self.source,self.root,'lecture-test')[0]
        box=[.08,.45,.35,.85]
        figure=add_figure(p,self.root,box)
        self.assertEqual(figure['rendering']['method'],'source-pdf-region')
        self.assertGreater(figure['width'],800)
        fid=figure['id'];p['figures']=[]
        self.assertEqual(add_figure(p,self.root,box)['id'],fid)

    def test_missing_source_uses_explicit_preview_fallback_without_upscaling(self):
        p=page();(self.root/'assets').mkdir();Image.new('RGB',(200,100),'white').save(self.root/p['image'])
        target,meta=render_region(p,self.root,[.1,.1,.4,.5],self.root/'crops')
        self.assertEqual(meta['method'],'preview-crop');self.assertEqual(meta['pixelSize'],[60,40])
        self.assertTrue(target.exists());self.assertEqual(meta['fallbackReason'],'source-pdf-unavailable')

    def test_cropped_rotated_pdf_coordinates_match_visible_source(self):
        original=self.root/'original.pdf';c=canvas.Canvas(str(original),pagesize=(600,400))
        c.setFont('Helvetica',15);c.drawString(170,200,'Visible source');c.drawString(10,20,'Outside crop');c.save()
        for rotation in (0,90,180,270):
            with self.subTest(rotation=rotation):
                writer=PdfWriter();pdf=PdfReader(original).pages[0]
                pdf.cropbox.lower_left=(100,100);pdf.cropbox.upper_right=(500,300);pdf.rotate(rotation)
                writer.add_page(pdf);writer.write(self.source)
                site=self.root/str(rotation);(site/'sources').mkdir(parents=True)
                (site/'sources/lecture-test.pdf').write_bytes(self.source.read_bytes())
                p=extract_pdf(self.source,site,'lecture-test')[0]
                self.assertNotIn('Outside',p['rawText'])
                # Use the union of glyph lines; rotated text may be tokenized vertically.
                lines=p['textLines'];self.assertTrue(lines)
                b=[min(l['position'][0] for l in lines),min(l['position'][1] for l in lines),max(l['position'][2] for l in lines),max(l['position'][3] for l in lines)]
                target,meta=render_region(p,site,b,site/'crops',padding=.005)
                with Image.open(target) as im:
                    dark=sum(im.convert('L').histogram()[:150])
                    self.assertGreater(dark,100)
                self.assertEqual(meta['rotation'],rotation)

    def test_one_detector_box_does_not_merge_three_columns(self):
        lines=[]
        for x in (.05,.37,.69):
            lines.extend([text_line('Question',[x,.2,x+.23,.23]),text_line('continued?',[x,.24,x+.2,.27])])
        blocks=native_blocks({'label':'text'},lines)
        self.assertEqual([b['text'] for b in blocks],['Question continued?']*3)
        self.assertTrue(all(len(b['sourceLines'])==2 for b in blocks))

    def test_paragraph_breaks_rules_font_changes_and_numbered_bullets(self):
        lines=[text_line('First paragraph',[.1,.2,.7,.23]),text_line('continued.',[.1,.24,.6,.27]),
               text_line('Second paragraph',[.1,.31,.7,.34]),text_line('Different style',[.1,.35,.7,.38],'Heading',18),
               text_line('1. First item',[.1,.42,.7,.45]),text_line('2. Second item',[.1,.46,.7,.49])]
        blocks=native_blocks({'label':'text'},lines)
        self.assertEqual([b['text'] for b in blocks],['First paragraph continued.','Second paragraph','Different style','1. First item','2. Second item'])
        self.assertEqual([b['type'] for b in blocks][-2:],['bullet','bullet'])
        separate=native_blocks({'label':'text'},lines[:2],[[.1,.235,.8,.235]])
        self.assertEqual(len(separate),2)

    def test_chinese_wraps_no_artificial_space_or_source_rewriting(self):
        blocks=native_blocks({'label':'text'},[text_line('需求价格',[.1,.2,.7,.23]),text_line('弹性',[.1,.24,.5,.27])])
        self.assertEqual(blocks[0]['text'],'需求价格弹性');self.assertEqual(blocks[0]['rawText'],'需求价格\n弹性')

    def test_native_merged_table_uses_real_cells_without_cloud_call(self):
        p=self.ruled_table();self.assertEqual(len(p['nativeTables']),1)
        table=p['nativeTables'][0];model=Mock()
        detected={'label':'table','bbox':table['bbox'],'score':.95,'readingOrder':0}
        with patch('course_compiler.local_layout.detect',return_value={'regions':[detected],'seconds':0,'model':'test','sha256':'test'}):
            result=parse(p,self.root/p['image'],model,{},self.root/'cache',output=self.root)
        model.request.assert_not_called();self.assertEqual(result['hybridMetrics']['nativeTables'],1)
        commit_layout(p,result,self.root)
        cells=[u for u in p['units'] if u['type']=='table-cell' and not u.get('reviewOnly')]
        self.assertEqual([c['sourceText'] for c in cells],['Total compensation','-6.25%','0.01'])
        self.assertEqual(cells[0]['colSpan'],2)
        self.assertAlmostEqual(cells[1]['position'][2],180/600)
        self.assertTrue(all(not c['positionEstimated'] for c in cells))
        self.assertEqual(p['tableRecognition'][0]['qualityEvidence']['method'],'native-grid-and-coverage')
        self.assertIsNone(p['tableRecognition'][0]['qualityEvidence']['textConfidence'])

    def test_invalid_exact_cell_coordinates_are_rejected(self):
        p=page()
        with self.assertRaisesRegex(Exception,'坐标'):
            commit_layout(p,{'tables':[{'bbox':[.1,.1,.9,.9],'cells':[{'row':0,'col':0,'text':'A','bbox':[0,0,.2,.2]}]}]},self.root)

    def test_scan_with_native_grid_never_bypasses_ocr(self):
        p=self.ruled_table();table=p['nativeTables'][0];p['imageRegions']=[table['bbox']]
        model=Mock();model.request.return_value='<table><tr><td>OCR</td></tr></table>'
        detected={'label':'table','bbox':table['bbox'],'score':.95,'readingOrder':0}
        with patch('course_compiler.local_layout.detect',return_value={'regions':[detected],'seconds':0,'model':'test','sha256':'test'}):
            result=parse(p,self.root/p['image'],model,{},self.root/'cache',output=self.root)
        self.assertEqual(model.request.call_count,1);self.assertEqual(result['hybridMetrics']['nativeTables'],0)
        commit_layout(p,result,self.root)
        self.assertEqual(p['tableRecognition'][0]['rendering']['method'],'source-pdf-region')

    def test_pptx_unequal_columns_and_merged_empty_cells_keep_physical_bounds(self):
        self.ruled_table();prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6])
        table=slide.shapes.add_table(2,2,Inches(1),Inches(1),Inches(6),Inches(2)).table
        table.columns[0].width=Inches(2);table.columns[1].width=Inches(4)
        table.rows[0].height=Inches(.5);table.rows[1].height=Inches(1.5)
        table.cell(0,0).merge(table.cell(0,1));table.cell(0,0).text='Merged'
        table.cell(1,0).text='Value'
        pptx=self.root/'lecture.pptx';prs.save(pptx)
        with patch('course_compiler.extract.convert_office',return_value=self.source):
            p=extract_pptx(pptx,pptx,self.root,'pptx')[0]
        cells=[u for u in p['units'] if u['type']=='table-cell']
        self.assertEqual(len(cells),3);self.assertEqual(cells[2]['sourceText'],'')
        self.assertEqual(cells[0]['colSpan'],2)
        self.assertAlmostEqual(cells[1]['position'][2],Inches(3)/prs.slide_width)
        self.assertAlmostEqual(cells[1]['position'][1],Inches(1.5)/prs.slide_height)
        self.assertEqual((self.root/'sources/pptx.pdf').read_bytes(),self.source.read_bytes())

    def test_old_completed_hybrid_cache_is_not_silently_recharged(self):
        source=self.root/'input.pdf';source.write_bytes(b'fixture');p=page()
        p.update(visualStatus='complete',transcriptionStatus='complete',validationMethod='local-layout-selective-ocr')
        p['units'][0].update(translatedText='译文',translationStatus='complete')
        profile={**PROFILE,'engine':'local-layout-ocr'}
        cid=course_id([source]);cache=self.root/'cache'/cid/(p['id']+'.json')
        atomic_json(cache,{'page':p,'stages':{'parse':'v5-adaptive-'+fingerprint(profile,hybrid_version='hybrid-1')}})
        parser=Mock();parser.parse.side_effect=AssertionError('Unexpected paid reparse')
        with patch('course_compiler.pipeline_v2.extract_files',return_value=[{'id':'lecture-test','name':'input.pdf','pages':[copy.deepcopy(p)]}]),patch('course_compiler.pipeline_v2.DocumentParser',return_value=parser),patch('course_compiler.local_layout.session'):
            c=compile_course([source],self.root/'site',{'parse':profile,'translation':PROFILE},self.root/'cache',workers=1,translation_context={'subject':'Test','style':'Faithful'})
        self.assertEqual(c['metrics']['documentCalls'],0);self.assertTrue(c['files'][0]['pages'][0]['parserUpgradeAvailable'])


if __name__=='__main__':unittest.main()
