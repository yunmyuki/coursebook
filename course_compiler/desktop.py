"""Coursebook desktop entry with native Windows and macOS windows."""
import argparse
import json
import os
import sys
import threading
import urllib.request
import webbrowser
from pathlib import Path
from .app_server import make_server
from .paths import data_root
from .storage import atomic_json,protect
from . import __version__

def main():
    parser=argparse.ArgumentParser()
    parser.add_argument('--self-test',metavar='REPORT')
    parser.add_argument('--test-office',action='store_true',help='Exercise PPTX and legacy PPT conversion during self-test')
    parser.add_argument('--window-test',metavar='REPORT')
    parser.add_argument('--browser',action='store_true')
    parser.add_argument('--serve-only',action='store_true',help='Run the loopback service without opening a window (automation)')
    parser.add_argument('--port',type=int,default=0)
    parser.add_argument('--data-dir')
    args=parser.parse_args()
    if args.data_dir:os.environ['COURSE_DATA_DIR']=str(Path(args.data_dir).resolve())
    root=data_root();root.mkdir(parents=True,exist_ok=True)
    diagnostic=None
    if args.self_test:
        import faulthandler
        diagnostic=(root/'self-test-trace.log').open('w')
        faulthandler.enable(diagnostic)
        faulthandler.dump_traceback_later(60,file=diagnostic)
    # Prevent two desktop windows from writing the same workspace simultaneously.
    handle=None
    workspace_lock=None
    if sys.platform=='darwin':
        import fcntl
        workspace_lock=(root/'desktop.lock').open('a')
        try:fcntl.flock(workspace_lock,fcntl.LOCK_EX|fcntl.LOCK_NB)
        except BlockingIOError:
            workspace_lock.close()
            return
    if os.name=='nt':
        import ctypes,hashlib
        api=ctypes.WinDLL('kernel32',use_last_error=True);api.CreateMutexW.restype=ctypes.c_void_p
        handle=api.CreateMutexW(None,False,'Local\\CourseCompiler-'+hashlib.sha256(str(root).encode()).hexdigest()[:16])
        if ctypes.get_last_error()==183:
            ctypes.windll.user32.MessageBoxW(None,'coursebook 已在运行。请切换到现有窗口。','coursebook',0)
            return
    server=make_server(args.port,root)
    thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start()
    url=f'http://127.0.0.1:{server.server_port}'
    try:
        if args.serve_only:
            thread.join()
            return
        if args.self_test:
            from .extract import extract_pdf,locate_converter
            from pypdf import PdfWriter
            from .parsers import unlimited_layout,glm_layout
            opener=urllib.request.build_opener(urllib.request.ProxyHandler({}))
            with opener.open(url+'/api/status') as response:status=json.load(response)
            with opener.open(url+'/') as response:html=response.read()
            with opener.open(url+'/web/desktop.js') as response:js=response.read()
            fixture=root/'self-test.pdf';writer=PdfWriter();writer.add_blank_page(width=300,height=200);writer.write(fixture)
            pages=extract_pdf(fixture,root/'self-test-site','lecture-selftest')
            result={'ok':status['version']==__version__ and b'coursebook' in html and len(js)==len((__import__('course_compiler.paths',fromlist=['resource_root']).resource_root()/'web/desktop.js').read_bytes()) and len(pages)==1,'frozen':bool(getattr(sys,'frozen',False)),'pdfRendered':True,'secureKeyRoundTrip':protect(protect('self-test-key'),True)=='self-test-key','officeAvailable':bool(locate_converter()),'version':status['version']}
            from .local_layout import detect,status as layout_status
            from PIL import Image,ImageDraw
            image=root/'layout-self-test.png';raster=Image.new('RGB',(1200,800),'white')
            draw=ImageDraw.Draw(raster);draw.text((80,100),'Coursebook release validation',fill='black',font_size=50)
            draw.text((80,250),'Preserve the full original paragraph.',fill='black',font_size=30)
            raster.save(image);layout=detect(image,root/'layout-cache',root)
            result['layoutReady']=layout_status(root)['ready']
            result['layoutInference']=bool(layout['regions'])
            result['layoutRegions']=len(layout['regions'])
            result['defaultHybrid']=status['settings']['profiles']['parse']['engine']=='local-layout-ocr'
            result['ok']=result['ok'] and result['layoutReady'] and result['layoutInference'] and result['defaultHybrid']
            if args.test_office:
                from pptx import Presentation
                from pptx.util import Inches
                from .extract import extract_pptx,convert_office
                prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6])
                slide.shapes.add_textbox(Inches(1),Inches(1),Inches(7),Inches(1)).text='Release test: 完整标题'
                table=slide.shapes.add_table(2,2,Inches(1),Inches(3),Inches(5),Inches(1)).table
                for i,text in enumerate(('Definition','Value','Gain','+5.0%')):table.cell(i//2,i%2).text=text
                slide.notes_slide.notes_text_frame.text='Preserve speaker notes.'
                pptx=root/'release-test.pptx';prs.save(pptx)
                extracted=extract_pptx(pptx,pptx,root/'pptx-site','lecture-release')
                units=extracted[0]['units']
                legacy=convert_office(pptx,root/'legacy','ppt')
                restored=convert_office(legacy,root/'roundtrip','pptx')
                result['officeRoundTrip']=restored.is_file() and len(extracted)==1 and sum(u['type']=='table-cell' for u in units)==4 and any(u['type']=='speaker-note' for u in units)
                result['ok']=result['ok'] and result['officeRoundTrip']
            import webview
            if sys.platform=='darwin':
                from webview.platforms import cocoa
            else:
                from webview.platforms import winforms
            result['webviewImported']=True
            atomic_json(Path(args.self_test),result)
            if not result['ok'] or not result['secureKeyRoundTrip']:raise RuntimeError('Release self-test failed')
            return
        if not args.browser:
            try:
                import webview
                webview.settings['ALLOW_DOWNLOADS']=True
                window=webview.create_window('Coursebook',url,width=1440,height=940,min_size=(900,650),background_color='#FFFFFF',hidden=bool(args.window_test))
                def closing():
                    if any(j['status']=='running' for j in server.jobs.values()):
                        if sys.platform=='darwin':
                            return window.create_confirmation_dialog('Coursebook','课程正在处理。关闭将暂停任务，已处理页面会保留。是否关闭？')
                        import ctypes
                        return ctypes.windll.user32.MessageBoxW(None,'课程正在处理。关闭将暂停任务，已处理页面会保留。是否关闭？','coursebook',0x24)==6
                    return True
                window.events.closing+=closing
                def smoke():
                    import time
                    window.events.loaded.wait(30)
                    for _ in range(40):
                        result=window.evaluate_js('({title:document.title,ready:!!document.getElementById("new-course"),loaded:!!document.querySelector(".empty-library,.course-row"),width:innerWidth})')
                        if result.get('loaded'):break
                        time.sleep(.25)
                    result['ok']=bool(result.get('ready') and result.get('loaded'));atomic_json(Path(args.window_test),result);window.destroy()
                webview.start(smoke if args.window_test else None,gui='cocoa' if sys.platform=='darwin' else 'edgechromium',private_mode=False,storage_path=str(root/'webview'))
                return
            except Exception:
                if args.window_test:raise
                # A usable local browser fallback also covers machines without WebView2.
                pass
        webbrowser.open(url)
        if os.name=='nt':
            import ctypes
            ctypes.windll.user32.MessageBoxW(None,'coursebook 已在浏览器中打开。\n处理课程时请保留此窗口。\n完成后点击“确定”退出。','coursebook',0)
        else:thread.join()
    except Exception as e:
        atomic_json(root/'startup-error.json',{'error':type(e).__name__,'message':str(e)})
        raise
    finally:
        for job in server.jobs.values():job['cancel'].set()
        server.shutdown();server.server_close()
        if workspace_lock:workspace_lock.close()
        if diagnostic:
            faulthandler.cancel_dump_traceback_later();faulthandler.disable();diagnostic.close()

if __name__=='__main__':main()
