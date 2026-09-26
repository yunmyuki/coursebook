"""Plan confirmed chapter insertion; compile new sources without rebuilding existing pages."""
import copy
import hashlib
import json
import re
import shutil
from datetime import datetime,timezone
from pathlib import Path
from .storage import atomic_json
from .pipeline import quality
from .pipeline_v2 import compile_course
from .export import export_site

def read_course(store,cid):
    path=store.course_path(cid)/'course.json'
    if not path.is_file():raise ValueError('课程不存在。')
    raw=path.read_bytes()
    return json.loads(raw),hashlib.sha256(raw).hexdigest()

def ordered_chapters(course):
    chapters=[{**ch,'fileId':f['id'],'fileName':f['name']} for f in course['files'] for ch in f['chapters']]
    by_id={ch['id']:ch for ch in chapters}
    order=course.get('readingOrder',[])
    if order and (len(order)!=len(by_id) or set(order)!=set(by_id)):raise ValueError('课程阅读顺序无效，请先修复。')
    return [by_id[c] for c in order] if order else chapters

def source_paths(store,selected):
    if not isinstance(selected,list) or not selected or len(selected)>100 or len(set(selected))!=len(selected):raise ValueError('请选择不重复的新材料。')
    paths=[]
    for fid in selected:
        if not isinstance(fid,str) or not re.fullmatch('[a-f0-9]{16}',fid):raise ValueError('文件编号无效。')
        found=[p for p in (store.root/'input'/fid).glob('*') if p.is_file() and p.suffix.lower() in ('.pdf','.pptx','.ppt')]
        if len(found)!=1:raise ValueError('材料不存在或同一文件有多个名称，请重新选择。')
        if hashlib.sha256(found[0].read_bytes()).hexdigest()[:16]!=fid:raise ValueError('材料已变化，请重新导入。')
        paths.append(found[0])
    return paths

def sample_text(path):
    """Read small native samples only; no rendering or paid recognition for the proposal."""
    try:
        if path.suffix.lower()=='.pdf':
            from pypdf import PdfReader
            reader=PdfReader(path)
            return ' '.join((p.extract_text() or '')[:1800] for p in reader.pages[:3])
        if path.suffix.lower()=='.pptx':
            from pptx import Presentation
            return ' '.join(s.text[:1000] for slide in list(Presentation(path).slides)[:3] for s in slide.shapes if s.has_text_frame)[:5400]
    except Exception:pass
    return ''

def words(text):
    stop={'the','and','for','with','from','this','that','lecture','chapter','course','part','week','notes','slide','overview','introduction','pdf','pptx'}
    return {w for w in re.findall(r'[a-z]{3,}|[\u3400-\u9fff]{2}',text.lower()) if w not in stop}

def sequence(text):
    match=re.search(r'(?:lecture|week|chapter|part|lec)[\s_.-]*(\d+)',text,re.I)
    return int(match.group(1)) if match else None

def plan_addition(store,cid,selected):
    course,revision=read_course(store,cid);chapters=ordered_chapters(course);paths=source_paths(store,selected)
    existing={f['sha256'] for f in course['files']}
    materials=[]
    for fid,path in zip(selected,paths):
        if hashlib.sha256(path.read_bytes()).hexdigest() in existing:raise ValueError('这份材料已在课程中，无需重复添加。')
        sample=sample_text(path);number=sequence(path.stem);before='';reason='未找到足够的顺序线索，建议放在课程末尾。';confidence='low'
        numbered=[(i,sequence(c['title']) or sequence(c['fileName'])) for i,c in enumerate(chapters)]
        later=next((i for i,n in numbered if n is not None and number is not None and n>number),None)
        if later is not None:before=chapters[later]['id'];reason='根据文件名中的课程编号，建议放在后续章节之前。';confidence='medium'
        else:
            tokens=words(path.stem+' '+sample)
            scores=[(len(tokens&words(c['title']))/max(1,len(words(c['title']))),i,len(tokens&words(c['title']))) for i,c in enumerate(chapters)]
            score,index,hits=max(scores,default=(0,0,0))
            if score>=.35 and hits>=2:
                before=chapters[index+1]['id'] if index+1<len(chapters) else ''
                reason='材料标题/文字与「'+chapters[index]['title']+'」相关，建议放在该章后。';confidence='medium'
        materials.append({'id':fid,'name':path.name,'beforeChapterId':before,'reason':reason,'confidence':confidence})
    return {'courseId':cid,'revision':revision,'materials':materials,'positions':[{'id':c['id'],'label':c['fileName']+' · '+c['title']} for c in chapters]}

def validate_addition(store,data):
    if data.get('extractOnly'):raise ValueError('追加材料需完成解析和翻译后加入，请先配置模型。')
    cid=data.get('courseId');course,revision=read_course(store,cid)
    if data.get('revision')!=revision:raise ValueError('课程已更新，请重新预览插入位置。')
    entries=data.get('materials',[])
    if not isinstance(entries,list) or any(not isinstance(m,dict) for m in entries):raise ValueError('材料位置格式无效。')
    paths=source_paths(store,[m.get('id') for m in entries]);valid={c['id'] for c in ordered_chapters(course)}|{''}
    if any(m.get('beforeChapterId') not in valid for m in entries):raise ValueError('请选择有效的章节位置。')
    existing={f['sha256'] for f in course['files']}
    if any(hashlib.sha256(p.read_bytes()).hexdigest() in existing for p in paths):raise ValueError('材料已在课程中，请重新预览。')
    return course,paths

def add_materials(store,data,settings,profiles,progress,cancel):
    original,paths=validate_addition(store,data);cid=original['id']
    digest=hashlib.sha256(''.join(sorted(m['id'] for m in data['materials'])).encode()).hexdigest()[:16]
    stage=store.root/'staging'/cid/digest
    context=original.get('translationContext')
    if context and context.get('status') not in ('generated','inherited'):context=None
    addition=compile_course(paths,stage,profiles,store.root/'cache',workers=settings['workers'],ai=not data.get('extractOnly'),explanations=False,parse_mode=settings['parseMode'],progress=progress,cancel=cancel,title=original['title'],request_limit=settings.get('requestLimit',1000),translation_context=context,terminology=original.get('terminology',{}))
    q=addition['quality']
    ready=not q['errors'] and (data.get('extractOnly') or (q['readingTranslated']==q['readingUnits'] and q['visualValidated']==q['pages'] and q['transcriptionValidated']==q['pages']))
    if not ready or addition['status']!='finished':return {'committed':False,'course':addition}
    with store.lock:
        current,revision=read_course(store,cid)
        if revision!=data['revision']:raise ValueError('处理期间课程已更新。新材料进度已保存，请重新预览后加入。')
        merged=copy.deepcopy(current);chapters=ordered_chapters(current)
        new_by_hash={f['sha256'][:16]:f for f in addition['files']}
        order=[]
        for before in [c['id'] for c in chapters]+['']:
            for entry in data['materials']:
                if entry['beforeChapterId']==before:order.extend(ch['id'] for ch in new_by_hash[entry['id']]['chapters'])
            if before:order.append(before)
        for entry in data['materials']:
            f=new_by_hash[entry['id']];f['order']=len(merged['files'])+1;merged['files'].append(f)
        merged['readingOrder']=order;merged['explanations'].extend(addition['explanations'])
        merged['terminology']={**addition.get('terminology',{}),**current.get('terminology',{})}
        if not context:merged['translationContext']=addition.get('translationContext',{})
        merged['compiledAt']=datetime.now(timezone.utc).isoformat();merged['quality']=quality(merged)
        if merged['quality']['errors']:raise ValueError('新材料关联校验失败，原课程未修改。')
        merged.setdefault('materialHistory',[]).append({'addedAt':merged['compiledAt'],'materials':data['materials']})
        site=store.course_path(cid)
        backup=store.root/'backups'/cid/('before-addition-'+revision[:16]+'.json');atomic_json(backup,current)
        for folder in ('assets','sources'):
            if (stage/folder).is_dir():shutil.copytree(stage/folder,site/folder,dirs_exist_ok=True)
        export_site(merged,site)
        for name,value in [('course.json',merged),('quality-report.json',merged['quality']),('terminology.json',merged['terminology'])]:atomic_json(site/name,value)
        item=next((c for c in store.courses() if c['id']==cid),{'id':cid,'title':merged['title']})
        mq=merged['quality'];complete=mq['readingTranslated']==mq['readingUnits'] and mq['visualValidated']==mq['pages'] and mq['transcriptionValidated']==mq['pages'] and not mq['errors']
        item.update(pages=mq['pages'],figures=mq.get('figures',0),quality=mq,updatedAt=merged['compiledAt'],status='partial' if not complete else 'review' if mq.get('reviewPages') else 'complete',readerUrl=f'/courses/{cid}/index.html',downloadUrl=f'/download/{cid}.zip')
        # Keep the original compile inputs intact; additional jobs have their own resume payload.
        item['hasAddedMaterials']=True;store.save_course(item)
        return {'committed':True,'course':merged,'metrics':addition['metrics']}
