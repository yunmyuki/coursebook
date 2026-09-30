"""Dual-path extraction: XML/text + rendered page; originals are never deduplicated."""
import hashlib
import re
import shutil
import subprocess
import unicodedata
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import pdfplumber
import pypdfium2 as pdfium
from .structure import restore_heading_groups
from pypdf.errors import PyPdfError
from pptx import Presentation


NS = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main', 'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart', 'm': 'http://schemas.openxmlformats.org/officeDocument/2006/math'}
BULLET = re.compile(r'^[\s]*([•●▪◦]|[–—]\s|\*\s|\d+[.)]\s)')


def natural_key(path):
    stem = Path(path).stem.lower()
    match = re.search(r'(?:lecture|week|chapter|part|lec)[\s_-]*(\d+)', stem)
    return (int(match.group(1)) if match else 10**9, [int(s) if s.isdigit() else s for s in re.split(r'(\d+)', stem)])


def course_order_key(path):
    """Use an explicit filename sequence, then a metadata sequence, then natural order."""
    path=Path(path)
    def sequence(text):
        named=re.search(r'(?:lecture|week|chapter|part|lec)[\s_.-]*(\d+)',text,re.I)
        number=named or re.search(r'(?<!\d)\d+(?!\d)',text)
        return int(named.group(1) if named else number.group()) if number else None
    number=sequence(path.stem)
    if number is None and path.is_file():
        try:
            metadata=''
            if path.suffix.lower()=='.pptx':
                with zipfile.ZipFile(path) as archive:
                    core=ET.fromstring(archive.read('docProps/core.xml'))
                    metadata=' '.join(el.text or '' for el in core if el.tag.rsplit('}',1)[-1] in ('title','subject','description'))
            elif path.suffix.lower()=='.pdf':
                from pypdf import PdfReader
                with path.open('rb') as stream:
                    if stream.read(5)!=b'%PDF':raise ValueError('No PDF metadata')
                info=PdfReader(str(path)).metadata or {}
                metadata=' '.join(str(info.get(k) or '') for k in ('/Title','/Subject'))
            # Generic numbers in metadata may be dates or software versions, not course order.
            match=re.search(r'(?:lecture|week|chapter|part|lec)[\s_.-]*(\d+)',metadata,re.I)
            if match:number=int(match.group(1))
        except (OSError,ValueError,KeyError,ET.ParseError,zipfile.BadZipFile,PyPdfError):
            pass
    return (number if number is not None else 10**9,natural_key(path)[1])


def normalize(text):
    # Compatibility normalization would turn x² into x2 and alter mathematical notation.
    return unicodedata.normalize('NFC', text).replace('\u00a0', ' ').strip()


def unit(page_id, i, text, kind, source, pos=None, **extra):
    return {'id': f'content-{page_id}-{i:04d}', 'type': kind, 'sourceText': normalize(text),
            'rawText': text, 'translatedText': None, 'translationStatus': 'pending',
            'source': dict(source), 'position': pos, 'level': 0, 'uncertain': False,
            'origin': 'text-layer', 'corrections': [], **extra}


def locate_converter():
    from .paths import resource_root,data_root
    return shutil.which('soffice') or next((str(p) for p in [resource_root()/'office/program/soffice.exe',data_root()/'components/office/program/soffice.exe',Path('C:/Program Files/LibreOffice/program/soffice.exe'), Path('/Applications/LibreOffice.app/Contents/MacOS/soffice')] if p.exists()), None)


def convert_office(path, destination, extension='pdf'):
    converter = locate_converter()
    if not converter:
        raise RuntimeError('PPT/PPTX 视觉渲染需要 LibreOffice。请安装后重试；原文件未改动。')
    destination.mkdir(parents=True, exist_ok=True)
    profile = (destination / 'lo-profile').resolve().as_uri()
    import os
    result = subprocess.run([converter, '-env:UserInstallation=' + profile, '--headless', '--convert-to', extension, '--outdir', str(destination), str(path.resolve())], capture_output=True, timeout=180,creationflags=subprocess.CREATE_NO_WINDOW if os.name=='nt' else 0)
    target = destination / (path.stem + '.' + extension)
    if result.returncode or not target.exists():
        raise RuntimeError('LibreOffice 转换失败；本文件未标记完成。')
    return target


def extract_pdf(path, output, file_id, progress=lambda *a: None, page_numbers=None):
    file_dir = output / 'assets' / file_id
    file_dir.mkdir(parents=True, exist_ok=True)
    pages = []
    renderer = pdfium.PdfDocument(str(path))
    with pdfplumber.open(path) as doc:
        for index, p in enumerate(doc.pages):
            if page_numbers is not None and index+1 not in page_numbers:continue
            page_id = f'{file_id}-page-{index+1:03d}'
            source = {'file': path.name, 'fileId': file_id, 'page': index+1}
            image_path = file_dir / f'page-{index+1:03d}.jpg'
            if not image_path.exists():
                render_page = renderer[index]
                bitmap = render_page.render(scale=min(2.3, 1700/max(p.width, p.height)))
                bitmap.to_pil().convert('RGB').save(image_path, quality=88)
                bitmap.close()
                render_page.close()
            from .layout import split_pdf_lines,group_pdf_lines
            text_lines,layout=split_pdf_lines(p)
            groups=group_pdf_lines(text_lines,layout,BULLET)
            units = []
            bullet_x = sorted(set(round(g['box'][0],2) for g in groups if g['kind']=='bullet'))
            for i,g in enumerate(groups,1):
                text = g['text']
                u = unit(page_id,i,text,g['kind'],source,g['box'],typography=g['style'],readingOrder=i-1)
                u['sourceText'] = re.sub(r'(?<=\w)-\n(?=[a-z])','',normalize(text)).replace('\n',' ')
                if u['sourceText'] != normalize(text):
                    u['corrections'].append({'kind':'line-reconstruction','before':text,'after':u['sourceText']})
                if g['kind']=='bullet':
                    u['level'] = min(4,bullet_x.index(round(g['box'][0],2)))
                units.append(u)
            raw_text = p.extract_text(y_tolerance=3) or ''
            links = []
            for link in p.hyperlinks:
                if link.get('uri'):
                    links.append({'url':link['uri'],'position':[link['x0']/p.width,link['top']/p.height,link['x1']/p.width,link['bottom']/p.height]})
            pages.append({'id':page_id,'number':index+1,'source':source,'title':units[0]['sourceText'] if units else f'Page {index+1}',
                          'layoutAnalysis':layout,'units':units,'rawText':raw_text,'rawUnits':[dict(u) for u in units],'textLines':text_lines,
                          'image':image_path.relative_to(output).as_posix(),'width':p.width,'height':p.height,
                          'links':links,'imageCount':len(p.images),'needsOCR':len(raw_text.strip())<50,
                          'imageRegions':[[max(0,im['x0']/p.width),max(0,im['top']/p.height),min(1,im['x1']/p.width),min(1,im['bottom']/p.height)] for im in p.images],
                          'vectorCount':len(p.curves)+len(p.lines)+len(p.rects),'tableCount':0,
                          'blankPage':not(p.chars or p.images or p.curves or p.lines or p.rects),
                          'visualStatus':'pending','transcriptionStatus':'pending','warnings':[]})
            restore_heading_groups(pages[-1])
            progress('extract',index+1,len(doc.pages),path.name)
    renderer.close()
    return pages


def extract_pptx(path, original, output, file_id, progress=lambda *a:None):
    """Use structured XML including tables/charts/notes and a required independent render."""
    converted = convert_office(path, output / '.intermediate' / file_id)
    pages = extract_pdf(converted,output,file_id,progress)
    prs = Presentation(str(path))
    if len(pages) != len(prs.slides):
        raise RuntimeError('PPT 渲染页数与 slide 数不一致，无法可靠关联。')
    for si, slide in enumerate(prs.slides):
        p = pages[si]
        p['source']['file'] = original.name
        p['units'] = []
        source = p['source']
        def add(text,kind,pos=None,**extra):
            if text.strip():
                p['units'].append(unit(p['id'],len(p['units'])+1,text,kind,source,pos,origin='pptx-xml',**extra))
        def walk(shapes,transform=(1,1,0,0)):
            for shape in sorted(shapes,key=lambda s:(s.top,s.left)):
                sx,sy,tx,ty=transform
                if hasattr(shape,'shapes'):
                    xfrm=shape._element.grpSpPr.xfrm
                    if xfrm is not None and xfrm.chExt is not None and xfrm.chExt.cx and xfrm.chExt.cy:
                        gx=shape.width/xfrm.chExt.cx;gy=shape.height/xfrm.chExt.cy
                        walk(shape.shapes,(sx*gx,sy*gy,tx+sx*(shape.left-gx*xfrm.chOff.x),ty+sy*(shape.top-gy*xfrm.chOff.y)))
                    else:walk(shape.shapes,transform)
                    continue
                box=[max(0,(sx*shape.left+tx)/prs.slide_width),max(0,(sy*shape.top+ty)/prs.slide_height),min(1,(sx*(shape.left+shape.width)+tx)/prs.slide_width),min(1,(sy*(shape.top+shape.height)+ty)/prs.slide_height)]
                if shape.has_table:
                    p['tableCount']=p.get('tableCount',0)+1
                    for row_i,row in enumerate(shape.table.rows):
                        for col_i,cell in enumerate(row.cells):
                            if cell.is_spanned:
                                continue
                            add(cell.text,'table-cell',box,tableId=f"{p['id']}-table-{shape.shape_id}",row=row_i,col=col_i,rowSpan=cell.span_height,colSpan=cell.span_width)
                elif shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        kind='title' if slide.shapes.title is not None and shape.shape_id==slide.shapes.title.shape_id else ('bullet' if para.level or para._p.find('.//'+ '{'+NS['a']+'}buChar') is not None else 'paragraph')
                        add(para.text,kind,box,level=para.level,hyperlinks=[r.hyperlink.address for r in para.runs if r.hyperlink.address],**({'headingContainerId':p['id']+'-shape-'+str(shape.shape_id)} if kind=='title' else {}))
                if shape.has_chart:
                    root=ET.fromstring(shape.chart._chartSpace.xml)
                    for el in root.findall('.//c:title//a:t',NS):
                        add(el.text or '', 'caption',box)
                    for cache in root.findall('.//c:strCache',NS):
                        for value in cache.findall('.//c:v',NS):
                            add(value.text or '', 'caption',box)
                    for series_i,series in enumerate(root.findall('.//c:ser',NS)):
                        for row_i,pt in enumerate(series.findall('.//c:numCache/c:pt',NS)):
                            value=pt.find('c:v',NS)
                            add(value.text if value is not None else '', 'table-cell',box,tableId=f"{p['id']}-chart-data-{shape.shape_id}",row=row_i,col=series_i)
                for math in ET.fromstring(shape._element.xml).findall('.//m:oMath',NS):
                    add(''.join(math.itertext()),'formula',box,mathXml=ET.tostring(math,encoding='unicode'),uncertain=True)
                # SmartArt and image text are supplied by the independent visual path.
        walk(slide.shapes)
        if slide.has_notes_slide:
            add(slide.notes_slide.notes_text_frame.text,'speaker-note')
        p['rawUnits']=[dict(u) for u in p['units']]
        p['rawText']='\n'.join(u['rawText'] for u in p['units'])
        p['title']=slide.shapes.title.text if slide.shapes.title else p['title']
        restore_heading_groups(p)
    return pages


def extract_files(paths, output, progress=lambda *a: None):
    files=[]
    used=set()
    for path in sorted([Path(p) for p in paths],key=course_order_key):
        digest=hashlib.sha256(path.read_bytes()).hexdigest()
        file_id='lecture-'+digest[:12]
        if file_id in used:
            raise ValueError('同一文件被重复选择；请保留一个副本。')
        used.add(file_id)
        dest=output/'sources'/f'{file_id}{path.suffix.lower()}'
        dest.parent.mkdir(parents=True,exist_ok=True)
        shutil.copy2(path,dest)
        if path.suffix.lower()=='.pdf':
            pages=extract_pdf(path,output,file_id,progress)
        else:
            converted=convert_office(path,output/'.intermediate'/file_id,'pptx') if path.suffix.lower()=='.ppt' else path
            pages=extract_pptx(converted,path,output,file_id,progress)
        files.append({'id':file_id,'name':path.name,'sha256':digest,'order':len(files)+1,'sourceUrl':dest.relative_to(output).as_posix(),'pages':pages})
    return files
